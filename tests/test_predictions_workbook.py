"""scripts/predictions_workbook.py: a day's predictions as a workbook. Synthetic rows, mechanics only."""
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from predict_bfsp_today import OUTPUT_COLS, with_model_rank
from scripts import predictions_workbook as pw


def _day():
    rng = np.random.default_rng(0)
    rows = []
    for t, track in [("1.30", "Newmarket"), ("2.05", "Haydock"), ("12.50", "Newmarket")]:
        p = rng.dirichlet(np.ones(5))
        for k in range(5):
            rows.append(dict(race_date="2026-09-25", race_time=t, track=track, race_name="H", horse_name=f"{track}{t}{k}",
                             predicted_bfsp=float(1 / p[k]), predicted_win_prob_norm=float(p[k]), odds=float(1.1 / p[k]),
                             raceid=f"{track}{t}", number_of_runners=5))
    return pd.DataFrame(rows)


def test_the_rank_and_the_columns():
    out = with_model_rank(_day())
    assert out.groupby("raceid")["model_rank"].min().eq(1).all()
    first = out.loc[out.groupby("raceid")["predicted_bfsp"].idxmin()]
    assert (first["model_rank"] == 1).all()
    assert {"model_rank", "predicted_bfsp", "jockey_name", "odds"} <= set(OUTPUT_COLS)


def test_the_workbook(tmp_path):
    path = pw.build(with_model_rank(_day()), {"feature_cols": ["x"] * 615, "best_iteration": 6000},
                    str(tmp_path / "p.xlsx"), note="test")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Read me", "Top picks", "All runners", "Newmarket", "Haydock"]
    top = wb["Top picks"]
    assert [top.cell(row=r, column=1).value for r in range(2, 5)] == ["12.50", "1.30", "2.05"]   # in running order
    assert str(top.cell(row=2, column=13).value).startswith("=IF(")                             # a formula, not a value
    allr = wb["All runners"]
    assert allr.max_row == 16 and str(allr.cell(row=2, column=allr.max_column).value).startswith("=IF(")


def test_the_rule_price_reads_the_one_margin_cell(tmp_path):
    path = pw.build(with_model_rank(_day()), {"feature_cols": ["x"] * 944}, str(tmp_path / "p.xlsx"), label="944")
    wb = load_workbook(path)
    readme = wb["Read me"]
    row = next(r for r in range(1, readme.max_row + 1) if readme.cell(row=r, column=1).value == "Margin (log)")
    assert readme.cell(row=row, column=2).value == pw.RULE_MARGIN
    ref = f"'Read me'!$B${row}"
    top = wb["Top picks"]
    assert top.cell(row=1, column=15).value == pw.RULE_HEAD
    assert top.cell(row=2, column=15).value == f'=IF(AND(ISNUMBER(H2),H2>0),H2*EXP({ref}),"")'
    allr = wb["All runners"]
    heads = [allr.cell(row=1, column=c).value for c in range(1, allr.max_column + 1)]
    p = heads.index("Predicted BSP") + 1
    col = heads.index(pw.RULE_HEAD) + 1
    letter = allr.cell(row=1, column=p).column_letter
    assert allr.cell(row=3, column=col).value == f'=IF(AND(ISNUMBER({letter}3),{letter}3>0),{letter}3*EXP({ref}),"")'


def test_the_comparison_sheet_picks_among_the_runners_still_declared(tmp_path):
    day = with_model_rank(_day())
    other = day[["track", "race_time", "horse_name", "predicted_bfsp"]].copy()
    other["predicted_bfsp"] = other["predicted_bfsp"][::-1].to_numpy()        # another model, other prices
    gone = other.groupby(["track", "race_time"])["predicted_bfsp"].idxmin()
    other = pd.concat([other, other.loc[gone].assign(horse_name="scratched", predicted_bfsp=1.01)])
    path = pw.build(day, {}, str(tmp_path / "p.xlsx"), label="944", other=other, other_label="615",
                    other_note="06:00")
    wb = load_workbook(path)
    assert wb.sheetnames[:4] == ["Read me", "Top picks", "944 vs 615", "All runners"]
    ws = wb["944 vs 615"]
    for r in range(2, 5):
        t, track = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
        race = day[(day["race_time"] == t) & (day["track"] == track)]
        o = other[(other["race_time"] == t) & (other["track"] == track) & (other["horse_name"] != "scratched")]
        assert ws.cell(row=r, column=5).value == race.loc[race["predicted_bfsp"].idxmin(), "horse_name"]
        assert ws.cell(row=r, column=8).value == o.loc[o["predicted_bfsp"].idxmin(), "horse_name"]
        assert ws.cell(row=r, column=10).value == f'=IF(E{r}=H{r},"yes","no")'
